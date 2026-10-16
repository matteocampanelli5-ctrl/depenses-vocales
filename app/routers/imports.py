"""Import de relevés — un routeur dédié (déviation mineure par rapport au
plan de découpage initial, qui ne listait pas de routeur "imports" à part :
jugé plus clair de séparer ce bloc dans son propre fichier plutôt que de le
noyer dans transactions.py/recurring.py).

Deux mécanismes complémentaires, valables pour n'importe quelle banque (pas
de format propriétaire codé en dur) :
  - import générique par IA : le contenu brut du fichier (.csv, .xlsx,
    .xlsm ou relevé .pdf) est envoyé à l'IA, qui le réinterprète elle-même
    et propose des transactions à valider. La détection récurrent / crédit /
    déjà-terminé, elle, reste déterministe côté serveur (voir
    _detect_recurring_patterns) ;
  - classeur onglet par onglet (/api/import/workbook/...) : pour un fichier
    Excel non standard, chaque onglet est reconnu (mensuel, prêt,
    endettement) et l'utilisateur confirme ce qu'il importe. Les prêts et
    l'historique de patrimoine sont lus en code (app/workbook_parsing.py) ;
    les onglets mensuels repassent par l'import générique, limité aux
    onglets choisis.
"""

import base64
import calendar
import json
import re
import unicodedata
from datetime import date
from io import BytesIO
from uuid import UUID

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile

from app.config import (
    ANTHROPIC_API_KEY,
    EXPENSE_CATEGORIES,
    INCOME_CATEGORIES,
    LOAN_CATEGORY,
    MAX_GENERIC_IMPORT_ENTRIES,
    TRANSFER_CATEGORY,
    TransactionType,
    category_label,
    today_paris,
)
from app.db import get_supabase_client
from app.emails import _notify_admin_once
from app.models import GenericImportCommitRequest, WorkbookCommitRequest
from app.routers.accounts import ensure_own_accounts
from app.routers.loans import active_loans, match_loan_for_payment
from app.routers.recurring import _clamp_day, _month_add
from app.security import require_user
from app.voice_nlp import _fetch_custom_categories
from app.workbook_parsing import analyze_workbook, property_key

router = APIRouter()


# ---------------------------------------------------------------------------
# Helpers partagés par l'import générique ci-dessous (_normalize_key sert
# aussi à la détection de récurrence : regrouper "iPad 3/12" et "iPad 4/12"
# malgré l'accentuation/la casse qui peut varier d'une banque à l'autre).
# ---------------------------------------------------------------------------
def _normalize_key(text: str) -> str:
    stripped = "".join(c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", stripped).strip().lower()


def _parse_french_amount(raw: str) -> float | None:
    cleaned = raw.strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    if not cleaned:
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None



# ---------------------------------------------------------------------------
# Import générique par IA (tout fichier .csv/.xlsx, structure quelconque,
# n'importe quelle banque) — le contenu brut du fichier est envoyé à l'IA qui
# le réinterprète elle-même. Utile par ex. pour un fichier personnel organisé
# en un tableau par mois, sans date précise par opération ni catégorie.
#
# La détection récurrent / crédit / déjà-terminé, elle, reste déterministe
# côté serveur (une comparaison de libellés/montants sur plusieurs mois est
# plus fiable en code qu'en laissant l'IA "se souvenir" d'un mois à l'autre) :
# voir _detect_recurring_patterns. Règles, dans l'ordre de priorité :
#   1. Libellé contenant un numéro d'échéance ("3/12", "échéance 4/18") :
#      quasi certain que c'est un crédit/financement fini, jamais une
#      dépense récurrente classique — le numéro permet même de calculer le
#      nombre d'échéances restantes et une date de fin.
#   2. Libellé contenant un nom d'organisme/abonnement connu (Netflix,
#      loyer, assurance...) : poids fort vers "récurrent", pas de
#      confirmation demandée.
#   3. Le motif a cessé d'apparaître dans les derniers mois disponibles du
#      fichier : considéré comme déjà terminé tout seul, pas de charge
#      récurrente créée (les occurrences passées sont importées telles
#      quelles).
#   4. Sinon (motif répété, ni mot-clé ni numéro d'échéance, toujours
#      présent récemment) : cas ambigu, l'IA ne tranche pas seule — bandeau
#      de confirmation côté frontend (même logique que l'assistant vocal).
# ---------------------------------------------------------------------------

MAX_GENERIC_IMPORT_FILE_SIZE = 3 * 1024 * 1024  # 3 Mo : largement assez pour un relevé annuel
# Relevé PDF : un peu plus large (mise en page, logos), mais sous la limite de
# 4,5 Mo par requête imposée par Vercel.
MAX_PDF_IMPORT_FILE_SIZE = 4 * 1024 * 1024
# Classeur onglet par onglet : un fichier Excel avec macros et beaucoup
# d'onglets est vite plus lourd ; même plafond Vercel.
MAX_WORKBOOK_FILE_SIZE = 4 * 1024 * 1024
MAX_GENERIC_IMPORT_TEXT_CHARS = 60_000  # ~15k tokens, large marge pour un relevé annuel

_SPREADSHEET_EXTENSIONS = (".xlsx", ".xlsm")
_ACCEPTED_EXTENSIONS = (".csv", ".pdf", *_SPREADSHEET_EXTENSIONS)

_KNOWN_RECURRING_KEYWORDS = {
    "netflix", "spotify", "disney", "canal+", "canal plus", "amazon prime",
    "deezer", "apple music", "apple one", "icloud", "appstore", "app store",
    "youtube premium", "loyer", "assurance", "mutuelle", "edf", "engie",
    "free mobile", "orange", "sfr", "bouygues", "box internet", "salle de sport",
    "abonnement", "creche", "cantine", "loa", "leasing",
}

# "3/12", "échéance 4/18", "mensualité 2 sur 10"...
_INSTALLMENT_PATTERN = re.compile(
    r"(?:(?:echeance|mensualite|prelevement)\s*)?(\d{1,2})\s*(?:/|sur)\s*(\d{1,2})\b",
    re.IGNORECASE,
)


def _strip_installment_counter(label: str) -> tuple[str, tuple[int, int] | None]:
    """Repère un numéro d'échéance dans un libellé brut et le retire pour ne
    garder que le nom (utilisé ensuite comme clé de regroupement commune à
    toutes les échéances d'un même crédit, ex: "iPad 3/12" et "iPad 4/12")."""
    normalized = _normalize_key(label)
    match = _INSTALLMENT_PATTERN.search(normalized)
    if not match:
        return label, None
    try:
        seen, total = int(match.group(1)), int(match.group(2))
    except ValueError:
        return label, None
    if seen < 1 or total < seen or total > 60:
        return label, None
    # Sur un relevé bancaire, "CB CARREFOUR 04/09" est la DATE du paiement
    # par carte, pas une échéance 4 sur 9 : sans mot-clé explicite
    # (échéance, mensualité...), un paiement par carte ou un "jj/mm" écrit
    # avec des zéros devant n'est jamais pris pour un numéro d'échéance.
    has_keyword = bool(re.search(r"echeance|mensualite|prelevement|\bsur\b", match.group(0)))
    compact = match.group(0).replace(" ", "")
    looks_like_date = bool(re.fullmatch(r"\d{2}/\d{2}", compact)) and 1 <= seen <= 31 and 1 <= total <= 12
    is_card_payment = bool(re.match(r"(cb|carte|paiement par carte)\b", normalized))
    if not has_keyword and (looks_like_date or is_card_payment):
        return label, None
    cleaned = (label[: match.start()] + label[match.end():]).strip(" -/:·|")
    return (cleaned or label), (seen, total)


def _yyyymm_to_tuple(value: str) -> tuple[int, int]:
    return int(value[:4]), int(value[5:7])


def _months_between(earlier: str, later: str) -> int:
    ey, em = _yyyymm_to_tuple(earlier)
    ly, lm = _yyyymm_to_tuple(later)
    return (ly * 12 + lm) - (ey * 12 + em)


def _add_months_to_date(d: date, months: int) -> date:
    year, month = _month_add(d.year, d.month, months)
    return date(year, month, _clamp_day(year, month, d.day))


def _extract_generic_import_text(filename: str, raw_bytes: bytes, only_sheets: list[str] | None = None) -> str:
    """Convertit le fichier (structure quelconque) en texte brut lisible par
    l'IA : un simple dump ligne par ligne, cellules séparées par " | ". Les
    lignes vides sont conservées (elles marquent souvent, dans ce genre de
    fichier personnel, la rupture entre deux tableaux mensuels).
    `only_sheets` : pour un classeur, ne garder que ces onglets (choisis par
    l'utilisateur à l'étape "onglet par onglet")."""
    if filename.endswith(_SPREADSHEET_EXTENSIONS):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(BytesIO(raw_bytes), data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(status_code=400, detail="Fichier Excel illisible") from exc
        lines: list[str] = []
        try:
            selected = [
                sheet for sheet in wb.worksheets
                if only_sheets is None or sheet.title in only_sheets
            ]
            if not selected:
                raise HTTPException(status_code=400, detail="Aucun des onglets choisis n'existe dans ce fichier")
            for sheet in selected:
                if len(wb.worksheets) > 1:
                    lines.append(f"### Feuille : {sheet.title}")
                if hasattr(sheet, "reset_dimensions"):
                    sheet.reset_dimensions()
                for row in sheet.iter_rows(max_row=3000, max_col=40, values_only=True):
                    cells = [str(c).strip() if c is not None else "" for c in row]
                    if not any(cells):
                        lines.append("")
                        continue
                    lines.append(" | ".join(cells))
        finally:
            wb.close()
        return "\n".join(lines)

    if filename.endswith(".csv"):
        try:
            return raw_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                return raw_bytes.decode("latin-1")
            except UnicodeDecodeError as exc:
                raise HTTPException(status_code=400, detail="Encodage du fichier non reconnu") from exc

    raise HTTPException(status_code=400, detail="Formats acceptés : .csv, .xlsx, .xlsm ou .pdf")


# Numéros de compte, IBAN, références bancaires : inutiles à l'IA pour
# reconnaître une opération, et sensibles. Masqués dans le texte envoyé
# (CSV / Excel) : IBAN, puis toute suite d'au moins 10 chiffres (éventuellement
# séparés par des espaces) — un montant ou une date n'en a jamais autant.
_IBAN_PATTERN = re.compile(r"\b[A-Z]{2}\d{2}(?:[ \u00a0]?[A-Z0-9]){11,30}\b")
_LONG_NUMBER_PATTERN = re.compile(r"(?<![\d.,])\d(?:[ \u00a0]?\d){9,}(?![\d])")


def _mask_sensitive_numbers(text: str) -> str:
    text = _IBAN_PATTERN.sub("[IBAN masqué]", text)
    return _LONG_NUMBER_PATTERN.sub("[n° masqué]", text)


_GENERIC_IMPORT_SYSTEM_PROMPT_TEMPLATE = """Tu reçois le contenu brut d'un fichier de suivi de finances personnelles (export Excel ou CSV), dont la structure n'est PAS standardisée — ce n'est pas forcément un relevé bancaire classique avec une colonne par champ. Il peut par exemple s'agir d'un tableau par mois (avec un titre de section indiquant le mois), sans date précise par opération, sans catégorie, juste un nom de dépense et un montant.

Ta tâche : repérer TOUTES les opérations (dépenses et revenus) présentes dans ce texte et les restituer sous forme d'une liste structurée.

Réponds UNIQUEMENT avec un objet JSON valide, sans aucun texte autour, selon ce schéma :
{
  "entries": [
    {
      "description": "libellé court et lisible de la dépense/du revenu (nettoyé, sans le montant)",
      "source_label": "le texte brut EXACT tel qu'il apparaît dans le fichier pour cette ligne (important : ne le modifie pas, il sert à détecter des motifs comme un numéro d'échéance)",
      "amount": nombre positif,
      "type": "expense" ou "income",
      "category_guess": une chaîne parmi __ALL_CATEGORIES__ — ta meilleure estimation à partir du libellé, "autre" si aucune ne correspond clairement,
      "month": "YYYY-MM" — le mois auquel se rattache cette opération, déduit du contexte (titre de section, en-tête de tableau, etc.),
      "day": nombre entier (jour du mois) UNIQUEMENT si une date précise est réellement indiquée pour CETTE opération dans le fichier, sinon null — n'invente JAMAIS un jour si le fichier ne donne qu'un mois,
      "transfer_account": si l'opération est un virement entre deux comptes de l'utilisateur lui-même (vers ou depuis un des comptes listés plus bas), le nom EXACT de cet autre compte tel qu'écrit dans la liste ; sinon null
    }
  ],
  "warnings": ["tout ce qui t'a semblé ambigu ou incertain lors de la lecture, en français" (liste vide si rien à signaler)]
}

RÈGLES IMPORTANTES :
- N'invente aucune opération : seulement celles réellement présentes dans le texte.
- Si l'année n'est pas explicitement indiquée, utilise __DEFAULT_YEAR__ par défaut et signale-le dans "warnings".
- Ignore les lignes de total, sous-total, solde, ou en-tête de colonnes.
- "source_label" doit rester le texte brut non modifié (utile pour une détection de motifs faite ensuite par un autre programme) ; "description" peut en revanche être nettoyé/raccourci.
- Conserve TOUTES les opérations, y compris celles qui se répètent d'un mois à l'autre (ne les regroupe pas toi-même) : un autre programme s'occupe ensuite de détecter les récurrences.
- Pour un relevé bancaire, une ligne au débit est une dépense ("expense") et une ligne au crédit un revenu ("income"), même pour un virement.

COMPTES DE L'UTILISATEUR (pour repérer les virements entre ses propres comptes, par exemple "VIR LIVRET A", "virement vers Bricks", "placement") : __ACCOUNTS__
Ne remplis "transfer_account" que si le libellé désigne clairement un de ces comptes ; dans le doute, null. Un virement vers ou depuis une autre personne n'est PAS un virement entre comptes de l'utilisateur."""


def _build_generic_import_prompt(all_categories: set[str], account_names: list[str]) -> str:
    accounts_text = (
        ", ".join(f'"{name}"' for name in account_names)
        if account_names
        else "aucun compte enregistré (mets toujours null pour transfer_account)"
    )
    return (
        _GENERIC_IMPORT_SYSTEM_PROMPT_TEMPLATE
        .replace("__ALL_CATEGORIES__", ", ".join(sorted(all_categories)))
        .replace("__DEFAULT_YEAR__", str(today_paris().year))
        .replace("__ACCOUNTS__", accounts_text)
    )


def _call_claude_generic_import(
    content: str | list[dict], all_categories: set[str], account_names: list[str]
) -> dict:
    if not ANTHROPIC_API_KEY:
        raise HTTPException(status_code=500, detail="Configuration manquante : ANTHROPIC_API_KEY")
    from anthropic import Anthropic

    client = Anthropic(api_key=ANTHROPIC_API_KEY)
    try:
        message = client.messages.create(
            model="claude-haiku-4-5-20251001",
            max_tokens=8000,
            system=_build_generic_import_prompt(all_categories, account_names),
            messages=[{"role": "user", "content": content}],
        )
    except Exception as exc:  # noqa: BLE001
        _notify_admin_once(
            "anthropic_api_failure",
            "⚠️ Kaching — l'IA ne répond plus",
            f"<p>L'appel à l'API Anthropic a échoué (plus de crédit, clé invalide, panne du "
            f"service...). Détail technique :</p><pre>{exc}</pre>",
        )
        raise HTTPException(
            status_code=502,
            detail="Le service d'import par IA est temporairement indisponible, réessaie plus tard.",
        ) from exc

    if getattr(message, "stop_reason", None) == "max_tokens":
        raise HTTPException(
            status_code=422,
            detail="Relevé trop long pour être lu en une fois : importe-le en plusieurs morceaux (un ou deux mois à la fois).",
        )
    raw = "".join(block.text for block in message.content if hasattr(block, "text")).strip()
    raw = re.sub(r"^```(json)?|```$", "", raw.strip(), flags=re.MULTILINE).strip()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise HTTPException(
            status_code=502, detail=f"Réponse IA invalide (JSON attendu) : {raw[:200]}"
        ) from exc
    if not isinstance(data, dict):
        raise HTTPException(status_code=502, detail="Réponse IA invalide (objet JSON attendu)")
    return data


def _detect_recurring_patterns(parsed: list[dict]) -> tuple[list[dict], dict[int, str]]:
    overall_last_month = max(row["expense_date"].strftime("%Y-%m") for row in parsed)

    groups: dict[tuple[str, str], list[int]] = {}
    group_counters: dict[tuple[str, str], list[tuple[int, int]]] = {}
    group_clean_label: dict[tuple[str, str], str] = {}

    for idx, row in enumerate(parsed):
        cleaned_label, counter = _strip_installment_counter(row["source_label"])
        # La date collée au libellé ("PRLV COFIDIS 05/10", "05/11"...) change
        # chaque mois : retirée de la clé de regroupement, sinon un même
        # prélèvement mensuel ne serait jamais reconnu comme récurrent.
        group_label = re.sub(r"\b\d{1,2}/\d{1,2}(?:/\d{2,4})?\b", " ", _normalize_key(cleaned_label))
        key = (row["type"], re.sub(r"\s+", " ", group_label).strip())
        groups.setdefault(key, []).append(idx)
        display_label = re.sub(r"\s+", " ", re.sub(r"\b\d{1,2}/\d{1,2}(?:/\d{2,4})?\b", " ", cleaned_label or "")).strip()
        group_clean_label.setdefault(key, (display_label or row["description"]).strip())
        if counter:
            group_counters.setdefault(key, []).append(counter)

    candidates: list[dict] = []
    group_by_index: dict[int, str] = {}
    group_num = 0

    for key, indices in groups.items():
        months_seen = sorted({parsed[i]["expense_date"].strftime("%Y-%m") for i in indices})
        has_counter = key in group_counters
        if len(months_seen) < 2 and not has_counter:
            continue  # occurrence isolée : pas assez de signal pour parler de récurrence

        group_num += 1
        group_id = f"g{group_num}"
        for i in indices:
            group_by_index[i] = group_id

        tx_type = key[0]
        amounts = sorted(parsed[i]["amount"] for i in indices)
        representative_amount = amounts[len(amounts) // 2]  # médiane : robuste aux variations ponctuelles
        categories = [parsed[i]["category"] for i in indices]
        category = max(set(categories), key=categories.count)
        first_month, last_month = months_seen[0], months_seen[-1]
        start_date = date(*_yyyymm_to_tuple(first_month), 1)

        installment_info: dict | None = None
        end_date: date | None = None

        if has_counter:
            seen_max, total_guess = max(group_counters[key], key=lambda c: c[0])
            remaining = max(total_guess - seen_max, 0)
            installment_info = {"last_seen": seen_max, "total": total_guess, "remaining": remaining}
            classification = "credit"
            needs_confirmation = False
            last_month_date = date(*_yyyymm_to_tuple(last_month), 1)
            end_date = _add_months_to_date(last_month_date, remaining) if remaining > 0 else last_month_date
        elif any(kw in _normalize_key(group_clean_label[key]) for kw in _KNOWN_RECURRING_KEYWORDS):
            classification = "recurring"
            needs_confirmation = False
        elif last_month != overall_last_month and _months_between(last_month, overall_last_month) >= 2:
            classification = "ended"
            needs_confirmation = False
            end_date = date(*_yyyymm_to_tuple(last_month), 1)
        else:
            classification = "uncertain"
            needs_confirmation = True

        candidates.append({
            "group_id": group_id,
            "name": group_clean_label[key][:80].strip().capitalize() or "Dépense récurrente",
            "type": tx_type,
            "amount": representative_amount,
            "category": category,
            "months_seen": months_seen,
            "classification": classification,
            "needs_confirmation": needs_confirmation,
            "installment_info": installment_info,
            "suggested_start_date": start_date,
            "suggested_end_date": end_date,
        })

    return candidates, group_by_index


def _parse_optional_uuid(raw: str | None) -> UUID | None:
    if not raw:
        return None
    try:
        return UUID(raw)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Compte invalide") from exc


def _parse_sheet_list(raw: str | None) -> list[str] | None:
    if not raw:
        return None
    try:
        sheets = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="Liste d'onglets invalide") from exc
    if not isinstance(sheets, list) or not all(isinstance(x, str) for x in sheets) or len(sheets) > 40:
        raise HTTPException(status_code=400, detail="Liste d'onglets invalide")
    return sheets


@router.post("/api/import/generic", response_model=None)
async def preview_generic_import(
    file: UploadFile = File(...),
    account_id: str | None = Form(None),
    sheets: str | None = Form(None),
    user_id: str = Depends(require_user),
) -> dict:
    filename = (file.filename or "").lower()
    if not filename.endswith(_ACCEPTED_EXTENSIONS):
        raise HTTPException(status_code=400, detail="Formats acceptés : .csv, .xlsx, .xlsm ou .pdf")
    is_pdf = filename.endswith(".pdf")

    raw_bytes = await file.read()
    if not raw_bytes:
        raise HTTPException(status_code=400, detail="Fichier vide")
    max_size = MAX_PDF_IMPORT_FILE_SIZE if is_pdf else (
        MAX_WORKBOOK_FILE_SIZE if filename.endswith(_SPREADSHEET_EXTENSIONS) else MAX_GENERIC_IMPORT_FILE_SIZE
    )
    if len(raw_bytes) > max_size:
        raise HTTPException(status_code=400, detail=f"Fichier trop volumineux ({max_size // (1024 * 1024)} Mo maximum)")

    if is_pdf:
        if not raw_bytes.startswith(b"%PDF"):
            raise HTTPException(status_code=400, detail="Ce fichier n'est pas un PDF valide")
        ai_content: str | list[dict] = [
            {
                "type": "document",
                "source": {
                    "type": "base64",
                    "media_type": "application/pdf",
                    "data": base64.b64encode(raw_bytes).decode("ascii"),
                },
            },
            {"type": "text", "text": "Voici le relevé à analyser."},
        ]
    else:
        raw_text = _mask_sensitive_numbers(_extract_generic_import_text(filename, raw_bytes, _parse_sheet_list(sheets)))
        if not raw_text.strip():
            raise HTTPException(status_code=400, detail="Fichier vide ou illisible")
        if len(raw_text) > MAX_GENERIC_IMPORT_TEXT_CHARS:
            raise HTTPException(
                status_code=422,
                detail=(
                    "Fichier trop volumineux/complexe pour l'import IA pour l'instant "
                    "(essaie un export sur une période plus courte, ou moins d'onglets à la fois)"
                ),
            )
        ai_content = raw_text

    client = get_supabase_client()
    statement_account = _parse_optional_uuid(account_id)
    ensure_own_accounts(client, user_id, statement_account)
    # Comptes vers/depuis lesquels un virement peut être reconnu : tous les
    # comptes actifs de l'utilisateur, sauf celui d'où vient le relevé.
    accounts = (
        client.table("accounts")
        .select("id, name, active")
        .eq("user_id", user_id)
        .execute()
        .data
    )
    transfer_targets = {
        _normalize_key(acc["name"]): acc["id"]
        for acc in accounts
        if acc.get("active", True) and acc["id"] != (str(statement_account) if statement_account else None)
    }
    target_names = [acc["name"] for acc in accounts if _normalize_key(acc["name"]) in transfer_targets]

    custom = _fetch_custom_categories(client, user_id)
    expense_cats = EXPENSE_CATEGORIES | set(custom.get("expense", {}).keys())
    income_cats = INCOME_CATEGORIES | set(custom.get("income", {}).keys())
    all_categories = expense_cats | income_cats
    loans = active_loans(client, user_id)

    ai_data = _call_claude_generic_import(ai_content, all_categories, target_names)
    raw_entries = ai_data.get("entries")
    if not isinstance(raw_entries, list) or not raw_entries:
        raise HTTPException(status_code=422, detail="Aucune opération exploitable détectée par l'IA dans ce fichier")
    truncated = len(raw_entries) > MAX_GENERIC_IMPORT_ENTRIES
    raw_entries = raw_entries[:MAX_GENERIC_IMPORT_ENTRIES]

    parsed: list[dict] = []
    skipped = 0
    for i, raw in enumerate(raw_entries):
        if not isinstance(raw, dict):
            skipped += 1
            continue
        amount_val = raw.get("amount")
        amount = _parse_french_amount(amount_val) if isinstance(amount_val, str) else None
        if amount is None and not isinstance(amount_val, str):
            try:
                amount = float(amount_val)
            except (TypeError, ValueError):
                amount = None
        if not amount:
            skipped += 1
            continue
        amount = round(abs(amount), 2)

        tx_type: TransactionType = "income" if raw.get("type") == "income" else "expense"

        month_str = str(raw.get("month") or "").strip()
        parts = month_str.split("-")
        try:
            if len(parts) < 2:
                raise ValueError
            year, month = int(parts[0]), int(parts[1])
            if not (1 <= month <= 12):
                raise ValueError
        except ValueError:
            skipped += 1
            continue

        day_raw = raw.get("day")
        day = None
        if isinstance(day_raw, (int, float)) and day_raw:
            candidate = int(day_raw)
            if 1 <= candidate <= calendar.monthrange(year, month)[1]:
                day = candidate
        expense_date = date(year, month, day or 1)

        description = str(raw.get("description") or "").strip() or "Opération importée"
        source_label = str(raw.get("source_label") or description).strip()
        category_guess = str(raw.get("category_guess") or "autre").strip().lower()
        valid_cats = income_cats if tx_type == "income" else expense_cats
        category = category_guess if category_guess in valid_cats else "autre"
        transfer_name = raw.get("transfer_account")
        transfer_account_id = (
            transfer_targets.get(_normalize_key(transfer_name)) if isinstance(transfer_name, str) else None
        )
        loan_id = None
        if transfer_account_id:
            category = TRANSFER_CATEGORY
        elif tx_type == "expense":
            # Mensualité d'un prêt de l'onglet Prêts (même montant, ou
            # libellé de prêt avec un montant proche) : rattachée à ce prêt.
            loan_id = match_loan_for_payment(
                loans, amount, f"{source_label} {description}", hinted=category == LOAN_CATEGORY, category=category
            )
            if loan_id:
                category = LOAN_CATEGORY

        parsed.append({
            "row_index": i,
            "expense_date": expense_date,
            "date_precision": "day" if day else "month",
            "type": tx_type,
            "amount": amount,
            "description": description,
            "category": category,
            "source_label": source_label,
            "transfer_account_id": transfer_account_id,
            "loan_id": loan_id,
        })

    if not parsed:
        raise HTTPException(status_code=422, detail="Aucune opération exploitable détectée par l'IA dans ce fichier")

    # Les virements entre comptes de l'utilisateur (ex: tous les mois vers le
    # Livret A) ne sont pas des charges récurrentes, et une mensualité de prêt
    # est déjà suivie dans l'onglet Prêts : exclus de la détection.
    flows = [row for row in parsed if not row["transfer_account_id"] and not row["loan_id"]]
    recurring_candidates, group_by_index = _detect_recurring_patterns(flows) if flows else ([], {})
    for row in parsed:
        row["recurring_group_id"] = None
    for idx, row in enumerate(flows):
        row["recurring_group_id"] = group_by_index.get(idx)

    dates = [row["expense_date"] for row in parsed]
    min_date, max_date = min(dates), max(dates)
    existing = (
        client.table("transactions")
        .select("expense_date, amount, type, description")
        .eq("user_id", user_id)
        .gte("expense_date", min_date.isoformat())
        .lte("expense_date", max_date.isoformat())
        .execute()
    ).data
    existing_signatures = {
        (
            row["expense_date"],
            round(float(row["amount"]), 2),
            row["type"],
            (row.get("description") or "").strip().lower(),
        )
        for row in existing
    }
    for row in parsed:
        sig = (row["expense_date"].isoformat(), row["amount"], row["type"], row["description"].strip().lower())
        row["likely_duplicate"] = sig in existing_signatures

    warnings = [w for w in (ai_data.get("warnings") or []) if isinstance(w, str)]
    if skipped:
        warnings.append(f"{skipped} ligne(s) ignorée(s) par l'IA car illisibles")
    if truncated:
        warnings.append(f"Fichier très volumineux : seules les {MAX_GENERIC_IMPORT_ENTRIES} premières opérations détectées ont été gardées")

    return {
        "rows": parsed,
        "recurring_candidates": recurring_candidates,
        "warnings": warnings,
        "expense_categories": {
            **{c: category_label(c, "expense") for c in EXPENSE_CATEGORIES},
            **custom.get("expense", {}),
        },
        "income_categories": {
            **{c: category_label(c, "income") for c in INCOME_CATEGORIES},
            **custom.get("income", {}),
        },
    }


@router.post("/api/import/generic/commit", status_code=201)
def commit_generic_import(req: GenericImportCommitRequest, user_id: str = Depends(require_user)) -> dict:
    if not req.transactions and not req.recurring:
        raise HTTPException(status_code=400, detail="Rien à importer")

    client = get_supabase_client()
    ensure_own_accounts(
        client, user_id, req.account_id, *(row.transfer_account_id for row in req.transactions)
    )
    account_id = str(req.account_id) if req.account_id else None
    inserted_transactions = 0
    inserted_recurring = 0
    loans = active_loans(client, user_id) if req.transactions else []
    own_loan_ids = {
        row["id"]
        for row in client.table("property_loans").select("id").eq("user_id", user_id).execute().data
    } if any(row.loan_id for row in req.transactions) else set()
    if any(row.loan_id and str(row.loan_id) not in own_loan_ids for row in req.transactions):
        raise HTTPException(status_code=400, detail="Prêt introuvable")

    if req.transactions:
        payload = []
        for row in req.transactions:
            # Un "virement" vers le compte du relevé lui-même n'en est pas un :
            # la ligne redevient une dépense/un revenu normal ("autre").
            transfer_id = (
                str(row.transfer_account_id)
                if row.transfer_account_id and str(row.transfer_account_id) != account_id
                else None
            )
            category = (row.category or "autre").strip() or "autre"
            loan_id = str(row.loan_id) if row.loan_id and not transfer_id else None
            if transfer_id:
                category = TRANSFER_CATEGORY
            elif category == TRANSFER_CATEGORY:
                category = "autre"
            if (
                not transfer_id
                and not loan_id
                and "loan_id" not in row.model_fields_set
                and row.type == "expense"
            ):
                # Ligne dont l'utilisateur n'a pas touché la catégorie : on
                # tente le rattachement à un prêt maintenant, ce qui couvre les
                # prêts créés juste avant par le même import de classeur.
                loan_id = match_loan_for_payment(
                    loans, row.amount, row.description, hinted=category == LOAN_CATEGORY, category=category
                )
            if loan_id:
                category = LOAN_CATEGORY
            payload.append({
                "type": row.type,
                "amount": row.amount,
                "category": category,
                "description": row.description,
                "expense_date": row.expense_date.isoformat(),
                "user_id": user_id,
                "account_id": account_id,
                "transfer_account_id": transfer_id,
                "loan_id": loan_id,
            })
        result = client.table("transactions").insert(payload).execute()
        inserted_transactions = len(result.data)

    if req.recurring:
        payload = [
            {
                "type": item.type,
                "name": item.name.strip(),
                "amount": item.amount,
                "category": (item.category or "autre").strip() or "autre",
                "day_of_month": item.day_of_month,
                "start_date": item.start_date.isoformat() if item.start_date else None,
                "end_date": item.end_date.isoformat() if item.end_date else None,
                "user_id": user_id,
                "active": True,
            }
            for item in req.recurring
        ]
        result = client.table("recurring_expenses").insert(payload).execute()
        inserted_recurring = len(result.data)

    return {"inserted_transactions": inserted_transactions, "inserted_recurring": inserted_recurring}


# ---------------------------------------------------------------------------
# Classeur onglet par onglet
# ---------------------------------------------------------------------------
@router.post("/api/import/workbook/analyze", response_model=None)
async def analyze_workbook_import(file: UploadFile = File(...), user_id: str = Depends(require_user)) -> dict:
    """Reconnaît le type de chaque onglet sans rien écrire en base ni appeler
    l'IA : l'utilisateur confirme ensuite onglet par onglet."""
    filename = (file.filename or "").lower()
    if not filename.endswith(_SPREADSHEET_EXTENSIONS):
        raise HTTPException(status_code=400, detail="Formats acceptés : .xlsx ou .xlsm")
    raw_bytes = await file.read()
    if not raw_bytes:
        raise HTTPException(status_code=400, detail="Fichier vide")
    if len(raw_bytes) > MAX_WORKBOOK_FILE_SIZE:
        raise HTTPException(status_code=400, detail="Fichier trop volumineux (4 Mo maximum)")
    try:
        sheets = analyze_workbook(raw_bytes, today_paris())
    except Exception as exc:  # noqa: BLE001 - fichier corrompu, protégé par mot de passe...
        raise HTTPException(status_code=400, detail="Fichier Excel illisible") from exc

    # Prêts déjà présents (même bien, même date de départ) : proposés
    # décochés, pour ne pas créer de doublon en réimportant le même fichier.
    client = get_supabase_client()
    existing = (
        client.table("property_loans")
        .select("property_name, start_date")
        .eq("user_id", user_id)
        .execute()
        .data
    )
    existing_keys = {(property_key(row["property_name"]), str(row["start_date"])[:10]) for row in existing}
    existing_names = {property_key(row["property_name"]) for row in existing}
    for sheet in sheets:
        loan = sheet.get("loan")
        if not loan:
            continue
        key = property_key(loan["property_name"])
        if (key, loan["start_date"]) in existing_keys:
            loan["already_imported"] = True
            sheet["selected"] = False
            sheet["note"] = "Déjà importé"
        else:
            loan["already_imported"] = False
            loan["same_name_exists"] = key in existing_names
    return {"sheets": sheets}


@router.post("/api/import/workbook/commit", status_code=201)
def commit_workbook_import(req: WorkbookCommitRequest, user_id: str = Depends(require_user)) -> dict:
    if not req.loans and not req.snapshots:
        raise HTTPException(status_code=400, detail="Rien à importer")
    client = get_supabase_client()
    inserted_loans = 0
    updated_rents = 0
    inserted_snapshots = 0

    if req.loans:
        existing = (
            client.table("property_loans")
            .select("id, property_name, start_date")
            .eq("user_id", user_id)
            .execute()
            .data
        )
        existing_by_key = {
            (property_key(row["property_name"]), str(row["start_date"])[:10]): row["id"] for row in existing
        }
        to_insert = []
        for loan in req.loans:
            key = (property_key(loan.property_name), loan.start_date.isoformat())
            if key in existing_by_key:
                # Même prêt déjà importé : on met juste le loyer à jour.
                if loan.monthly_rent is not None:
                    client.table("property_loans").update({"monthly_rent": loan.monthly_rent}).eq(
                        "id", existing_by_key[key]
                    ).eq("user_id", user_id).execute()
                    updated_rents += 1
                continue
            to_insert.append({
                "user_id": user_id,
                "property_name": loan.property_name.strip(),
                "principal_amount": loan.principal_amount,
                "annual_rate_pct": loan.annual_rate_pct,
                "monthly_payment": loan.monthly_payment,
                "start_date": loan.start_date.isoformat(),
                "monthly_rent": loan.monthly_rent,
                "active": loan.active,
            })
        if to_insert:
            inserted_loans = len(client.table("property_loans").insert(to_insert).execute().data)

    if req.snapshots:
        existing_dates = {
            str(row["snapshot_date"])[:10]
            for row in client.table("net_worth_snapshots")
            .select("snapshot_date")
            .eq("user_id", user_id)
            .execute()
            .data
        }
        rows = []
        for snap in req.snapshots:
            day = snap.snapshot_date.isoformat()
            if day in existing_dates:
                continue
            existing_dates.add(day)
            rows.append({
                "user_id": user_id,
                "snapshot_date": day,
                "assets": round(snap.assets, 2),
                "debts": round(snap.debts, 2),
                "label": (snap.label or "").strip() or None,
            })
        if rows:
            inserted_snapshots = len(client.table("net_worth_snapshots").insert(rows).execute().data)

    return {
        "inserted_loans": inserted_loans,
        "updated_rents": updated_rents,
        "inserted_snapshots": inserted_snapshots,
    }
