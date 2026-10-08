"""Export JSON (sauvegarde complète) et Excel (.xlsx)."""

import json
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, Response

from app.config import _RESET_TABLES, category_label, today_paris
from app.db import get_supabase_client
from app.routers.recurring import sync_recurring_occurrences
from app.security import require_user

router = APIRouter()


# ---------------------------------------------------------------------------
# Sauvegarde complète (JSON brut de toutes les tables) — Supabase ne fait pas
# de sauvegarde automatique en offre gratuite ; ce fichier est le filet de
# sécurité en cas de suppression accidentelle ou de pépin côté Supabase. Ne
# couvre pas les photos de reçus elles-mêmes (binaires), seulement les
# chemins qui y pointent (receipt_path) — les fichiers restent dans le bucket
# Storage, qui a sa propre durabilité côté Supabase.
# ---------------------------------------------------------------------------
@router.get("/api/export/json")
def export_json(user_id: str = Depends(require_user)) -> Response:
    client = get_supabase_client()
    sync_recurring_occurrences(client, user_id)

    backup: dict = {
        "app": "Kaching",
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }
    for table_name in _RESET_TABLES:
        backup[table_name] = client.table(table_name).select("*").eq("user_id", user_id).execute().data

    payload = json.dumps(backup, indent=2, ensure_ascii=False, default=str)
    filename = f"kaching-sauvegarde-{today_paris().isoformat()}.json"
    return Response(
        content=payload,
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Export Excel
# ---------------------------------------------------------------------------
@router.get("/api/export/xlsx")
def export_xlsx(user_id: str = Depends(require_user)) -> Response:
    client = get_supabase_client()
    sync_recurring_occurrences(client, user_id)
    transactions = (
        client.table("transactions")
        .select("*")
        .eq("user_id", user_id)
        .order("expense_date", desc=True)
        .order("created_at", desc=True)
        .execute()
    ).data
    recurring = (
        client.table("recurring_expenses")
        .select("*")
        .eq("user_id", user_id)
        .order("day_of_month")
        .execute()
    ).data
    category_ratings = (
        client.table("category_essentiality_ratings")
        .select("*")
        .eq("user_id", user_id)
        .execute()
    ).data

    # Import différé : openpyxl n'est utile que pour cette seule route, pas
    # besoin de l'importer (et de ralentir le démarrage de la fonction) pour
    # le reste de l'API.
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.series import DataPoint
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.styles import Font, PatternFill

    type_labels = {"expense": "Dépense", "income": "Revenu"}

    # Mêmes couleurs que le site (voir CHART_COLORS côté frontend) pour que
    # l'export ressemble au tableau de bord : rouge pour les dépenses, vert
    # pour les revenus, et la même palette cyclique pour les catégories.
    COLOR_EXPENSE = "EF4444"
    COLOR_INCOME = "22C55E"
    COLOR_ACCENT = "3B82F6"
    CHART_COLORS = ["3B82F6", "22C55E", "EF4444", "F59E0B", "A855F7", "14B8A6", "EC4899", "64748B"]
    HEADER_FILL = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

    def category_colors(rows: list[tuple[str, float]]) -> list[DataPoint]:
        return [
            DataPoint(idx=i, spPr=GraphicalProperties(solidFill=CHART_COLORS[i % len(CHART_COLORS)]))
            for i in range(len(rows))
        ]

    # --- Agrégats pour le résumé (totaux, camemberts, évolution mensuelle) ---
    total_expenses = 0.0
    total_income = 0.0
    expense_category_totals: dict[str, float] = {}
    income_category_totals: dict[str, float] = {}
    monthly_totals: dict[str, dict[str, float]] = {}
    for tx in transactions:
        amount = float(tx["amount"])
        month_key = tx["expense_date"][:7]
        monthly_totals.setdefault(month_key, {"expense": 0.0, "income": 0.0})
        monthly_totals[month_key][tx["type"]] += amount
        if tx["type"] == "expense":
            total_expenses += amount
            expense_category_totals[tx["category"]] = expense_category_totals.get(tx["category"], 0.0) + amount
        else:
            total_income += amount
            income_category_totals[tx["category"]] = income_category_totals.get(tx["category"], 0.0) + amount
    expense_category_rows = sorted(expense_category_totals.items(), key=lambda kv: kv[1], reverse=True)
    income_category_rows = sorted(income_category_totals.items(), key=lambda kv: kv[1], reverse=True)
    month_keys = sorted(monthly_totals)
    solde = total_income - total_expenses

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Résumé"
    ws_summary["A1"] = "Résumé financier"
    ws_summary["A1"].font = Font(bold=True, size=14, color=COLOR_ACCENT)

    ws_summary["A3"] = "Solde"
    ws_summary["B3"] = solde
    ws_summary["A4"] = "Total dépenses"
    ws_summary["B4"] = total_expenses
    ws_summary["A5"] = "Total revenus"
    ws_summary["B5"] = total_income
    ws_summary.cell(row=3, column=1).font = Font(bold=True)
    ws_summary.cell(row=3, column=2).font = Font(bold=True, color=COLOR_INCOME if solde >= 0 else COLOR_EXPENSE)
    ws_summary.cell(row=4, column=1).font = Font(bold=True, color=COLOR_EXPENSE)
    ws_summary.cell(row=4, column=2).font = Font(color=COLOR_EXPENSE)
    ws_summary.cell(row=5, column=1).font = Font(bold=True, color=COLOR_INCOME)
    ws_summary.cell(row=5, column=2).font = Font(color=COLOR_INCOME)

    def write_category_section(title: str, header_row: int, rows: list[tuple[str, float]], tx_type: str) -> int:
        """Écrit le titre de section + le tableau Catégorie/Montant à partir de
        `header_row`, et renvoie la dernière ligne utilisée (ou header_row si
        `rows` est vide, pour que l'appelant puisse quand même avancer)."""
        ws_summary.cell(row=header_row - 1, column=1, value=title).font = Font(bold=True, color=COLOR_ACCENT)
        ws_summary.cell(row=header_row, column=1, value="Catégorie")
        ws_summary.cell(row=header_row, column=2, value="Montant (€)")
        for col in (1, 2):
            cell = ws_summary.cell(row=header_row, column=col)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
        start_row = header_row + 1
        for i, (cat, amount) in enumerate(rows):
            row = start_row + i
            ws_summary.cell(row=row, column=1, value=category_label(cat, tx_type))
            ws_summary.cell(row=row, column=2, value=amount)
        return start_row + len(rows) - 1 if rows else header_row

    # Les deux camemberts sont anchorés à intervalle fixe en colonne D, assez
    # espacés (17 lignes) pour ne jamais se chevaucher quel que soit le
    # nombre de catégories listées en colonnes A/B à côté.
    cat_header_row = 7
    cat_end_row = write_category_section(
        "Dépenses par catégorie", cat_header_row, expense_category_rows, "expense"
    )
    if expense_category_rows:
        pie_expense = PieChart()
        pie_expense.title = "Répartition des dépenses par catégorie"
        data = Reference(ws_summary, min_col=2, min_row=cat_header_row + 1, max_row=cat_end_row)
        cats = Reference(ws_summary, min_col=1, min_row=cat_header_row + 1, max_row=cat_end_row)
        pie_expense.add_data(data, titles_from_data=False)
        pie_expense.set_categories(cats)
        pie_expense.series[0].data_points = category_colors(expense_category_rows)
        pie_expense.height = 8
        pie_expense.width = 13
        ws_summary.add_chart(pie_expense, "D3")

    inc_header_row = cat_end_row + 3
    inc_end_row = write_category_section(
        "Revenus par catégorie", inc_header_row, income_category_rows, "income"
    )
    if income_category_rows:
        pie_income = PieChart()
        pie_income.title = "Répartition des revenus par catégorie"
        data = Reference(ws_summary, min_col=2, min_row=inc_header_row + 1, max_row=inc_end_row)
        cats = Reference(ws_summary, min_col=1, min_row=inc_header_row + 1, max_row=inc_end_row)
        pie_income.add_data(data, titles_from_data=False)
        pie_income.set_categories(cats)
        pie_income.series[0].data_points = category_colors(income_category_rows)
        pie_income.height = 8
        pie_income.width = 13
        ws_summary.add_chart(pie_income, "D20")

    month_header_row = inc_end_row + 3
    ws_summary.cell(row=month_header_row - 1, column=1, value="Évolution mensuelle").font = Font(
        bold=True, color=COLOR_ACCENT
    )
    ws_summary.cell(row=month_header_row, column=1, value="Mois")
    ws_summary.cell(row=month_header_row, column=2, value="Dépenses")
    ws_summary.cell(row=month_header_row, column=3, value="Revenus")
    for col in (1, 2, 3):
        cell = ws_summary.cell(row=month_header_row, column=col)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    month_start_row = month_header_row + 1
    for i, month_key in enumerate(month_keys):
        row = month_start_row + i
        ws_summary.cell(row=row, column=1, value=month_key)
        ws_summary.cell(row=row, column=2, value=monthly_totals[month_key]["expense"]).font = Font(color=COLOR_EXPENSE)
        ws_summary.cell(row=row, column=3, value=monthly_totals[month_key]["income"]).font = Font(color=COLOR_INCOME)
    month_end_row = month_start_row + len(month_keys) - 1

    if month_keys:
        bar = BarChart()
        bar.type = "col"
        bar.title = "Évolution mensuelle (dépenses vs revenus)"
        data = Reference(ws_summary, min_col=2, max_col=3, min_row=month_header_row, max_row=month_end_row)
        cats = Reference(ws_summary, min_col=1, min_row=month_start_row, max_row=month_end_row)
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(cats)
        bar.series[0].graphicalProperties.solidFill = COLOR_EXPENSE
        bar.series[1].graphicalProperties.solidFill = COLOR_INCOME
        bar.height = 8
        bar.width = 15
        ws_summary.add_chart(bar, "D37")

    for col_letter, width in zip("ABC", [24, 14, 14]):
        ws_summary.column_dimensions[col_letter].width = width

    # --- Transactions : deux blocs de colonnes séparés (dépenses / revenus) ---
    # plutôt qu'une seule liste mélangée triée par date — plus lisible, en
    # particulier avec beaucoup de lignes des deux types.
    ws_tx = wb.create_sheet("Transactions")
    expense_txs = [tx for tx in transactions if tx["type"] == "expense"]
    income_txs = [tx for tx in transactions if tx["type"] == "income"]

    def write_tx_block(title: str, start_col: int, rows: list[dict], color: str) -> None:
        col_letters = [chr(ord("A") + start_col - 1 + i) for i in range(4)]
        ws_tx.cell(row=1, column=start_col, value=title).font = Font(bold=True, color=COLOR_ACCENT)
        headers = ["Date", "Catégorie", "Montant (€)", "Description"]
        for i, header in enumerate(headers):
            cell = ws_tx.cell(row=2, column=start_col + i, value=header)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
        for r, tx in enumerate(rows):
            row = 3 + r
            ws_tx.cell(row=row, column=start_col, value=tx["expense_date"])
            ws_tx.cell(row=row, column=start_col + 1, value=category_label(tx["category"], tx["type"]))
            ws_tx.cell(row=row, column=start_col + 2, value=float(tx["amount"])).font = Font(color=color)
            ws_tx.cell(row=row, column=start_col + 3, value=tx.get("description") or "")
        for col_letter, width in zip(col_letters, [12, 14, 12, 42]):
            ws_tx.column_dimensions[col_letter].width = width

    write_tx_block("Dépenses", 1, expense_txs, COLOR_EXPENSE)  # colonnes A-D
    write_tx_block("Revenus", 6, income_txs, COLOR_INCOME)  # colonnes F-I (E = séparateur)

    ws_rec = wb.create_sheet("Récurrentes")
    ws_rec.append([
        "Type", "Nom", "Montant (€)", "Catégorie", "Jour du mois", "Date de début", "Date de fin",
        "Essentialité (1-5)",
    ])
    for cell in ws_rec[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for item in recurring:
        rec_type = item.get("type") or "expense"
        ws_rec.append([
            type_labels.get(rec_type, rec_type),
            item["name"],
            float(item["amount"]),
            category_label(item["category"], rec_type),
            item["day_of_month"],
            item.get("start_date") or "",
            item.get("end_date") or "",
            item.get("essentiality_rating") or "",
        ])
        ws_rec.cell(row=ws_rec.max_row, column=3).font = Font(
            color=COLOR_INCOME if rec_type == "income" else COLOR_EXPENSE
        )
    for col_letter, width in zip("ABCDEFGH", [10, 24, 12, 14, 12, 14, 14, 16]):
        ws_rec.column_dimensions[col_letter].width = width

    # Plan d'épargne : notes d'essentialité par catégorie (celles des charges
    # récurrentes sont déjà dans la colonne ci-dessus) — uniquement si
    # l'utilisateur en a renseigné au moins une.
    if category_ratings:
        ws_plan = wb.create_sheet("Plan d'épargne")
        ws_plan.append(["Catégorie", "Essentialité (1-5)"])
        for cell in ws_plan[1]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
        for rating_row in sorted(category_ratings, key=lambda r: r["category"]):
            ws_plan.append([
                category_label(rating_row["category"], "expense"),
                rating_row["rating"],
            ])
        for col_letter, width in zip("AB", [24, 16]):
            ws_plan.column_dimensions[col_letter].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"depenses_{today_paris().isoformat()}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
