"""Lecture d'un classeur Excel non standard, onglet par onglet, sans IA.

Pensé pour des fichiers personnels de suivi de patrimoine : un onglet par
mois (dépenses/revenus en vrac), un onglet par prêt
(tableau d'amortissement), un onglet "Endettement" avec plusieurs tableaux
côte à côte. Ce module ne fait que LIRE et PROPOSER : il reconnaît le type
probable de chaque onglet et extrait ce qui peut l'être de façon sûre (les
prêts et l'historique de patrimoine sont des tableaux réguliers, lus en
code). Les onglets mensuels, eux, sont trop libres pour être lus en code :
ils partent à l'IA via l'import générique (voir app/routers/imports.py).

Rien n'est écrit en base ici, et aucune donnée n'est conservée : tout est
renvoyé au navigateur, qui demande confirmation onglet par onglet.

Aucune dépendance à FastAPI/Supabase : seulement openpyxl, pour pouvoir
tester ce module directement sur un vrai fichier.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from io import BytesIO

# Bornes de lecture : un onglet de suivi personnel dépasse rarement quelques
# centaines de lignes ; au-delà on s'arrête plutôt que de consommer tout le
# temps d'exécution de la fonction serverless sur un onglet de données brutes.
MAX_SCAN_ROWS = 1500
MAX_SCAN_COLS = 40
MAX_SHEETS = 40

SHEET_KINDS = ("monthly", "loan", "net_worth", "ignore")

_MONTHS_FR = {
    "janvier": 1, "janv": 1, "jan": 1,
    "fevrier": 2, "fevr": 2, "fev": 2,
    "mars": 3, "mar": 3,
    "avril": 4, "avr": 4,
    "mai": 5,
    "juin": 6,
    "juillet": 7, "juil": 7,
    "aout": 8,
    "septembre": 9, "sept": 9, "sep": 9,
    "octobre": 10, "oct": 10,
    "novembre": 11, "nov": 11,
    "decembre": 12, "dec": 12,
}


def norm(value) -> str:
    """Texte normalisé pour comparer des libellés : sans accents, en
    minuscules, espaces multiples réduits."""
    if value is None:
        return ""
    text = str(value)
    text = "".join(c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", text).strip().lower()


def property_key(name: str) -> str:
    """Clé de rapprochement entre deux écritures du même bien ("La Résidence"
    dans un onglet, "Le Résidence" dans un autre) : sans article ni
    ponctuation."""
    text = norm(name)
    text = re.sub(r"^(?:(?:les|le|la)\s+|l'\s*)", "", text)
    return re.sub(r"[^a-z0-9]", "", text)


def sheet_month(name: str) -> str | None:
    """'oct 2026', 'Octobre 2026', '2026-10', '10/2026' -> '2026-10'."""
    text = norm(name)
    match = re.search(r"([a-z]+)\.?[\s\-_/]*((?:19|20)\d{2})", text)
    if match and match.group(1) in _MONTHS_FR:
        return f"{match.group(2)}-{_MONTHS_FR[match.group(1)]:02d}"
    match = re.search(r"((?:19|20)\d{2})[\s\-_/.]+(\d{1,2})\b", text)
    if match and 1 <= int(match.group(2)) <= 12:
        return f"{match.group(1)}-{int(match.group(2)):02d}"
    match = re.search(r"\b(\d{1,2})[\s\-_/.]+((?:19|20)\d{2})", text)
    if match and 1 <= int(match.group(1)) <= 12:
        return f"{match.group(2)}-{int(match.group(1)):02d}"
    return None


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _as_number(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        cleaned = value.replace("\xa0", "").replace(" ", "").replace("€", "").replace(",", ".")
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def _is_month_or_year_label(value) -> bool:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return 1900 <= value <= 2100
    text = norm(value)
    return text in _MONTHS_FR or bool(re.fullmatch(r"(19|20)\d{2}", text))


def _grid(ws) -> list[list]:
    # Classeur ouvert en lecture seule (beaucoup plus rapide : ~0,3 s au lieu
    # de ~12 s sur un fichier de 800 Ko avec macros). Les dimensions
    # déclarées dans le fichier sont ignorées (parfois fausses selon le
    # logiciel qui l'a enregistré) : on lit jusqu'aux bornes ci-dessus.
    if hasattr(ws, "reset_dimensions"):
        ws.reset_dimensions()
    rows: list[list] = []
    for row in ws.iter_rows(min_row=1, max_row=MAX_SCAN_ROWS, max_col=MAX_SCAN_COLS, values_only=True):
        rows.append(list(row))
    while rows and not any(v not in (None, "") for v in rows[-1]):
        rows.pop()
    return rows


def _cell(rows: list[list], r: int, c: int):
    if 0 <= r < len(rows) and 0 <= c < len(rows[r]):
        return rows[r][c]
    return None


def _find_header_rows(rows: list[list], required: dict[str, tuple[str, ...]]) -> list[tuple[int, dict[str, int]]]:
    """Lignes dont les cellules contiennent tous les mots-clés demandés.
    `required` : nom logique -> sous-chaînes acceptées (normalisées). Renvoie
    (index de ligne, nom logique -> index de colonne) ; la première colonne
    trouvée à gauche gagne quand un mot-clé apparaît plusieurs fois (deux
    tableaux côte à côte sur la même ligne)."""
    found: list[tuple[int, dict[str, int]]] = []
    for r, row in enumerate(rows):
        texts = [norm(v) if isinstance(v, str) else "" for v in row]
        if not any(texts):
            continue
        # Plusieurs tableaux peuvent partager la même ligne d'en-tête : on
        # cherche chaque occurrence du premier mot-clé et on regarde les
        # colonnes qui suivent immédiatement.
        first_key = next(iter(required))
        for start, text in enumerate(texts):
            if not any(k in text for k in required[first_key]):
                continue
            colmap = {first_key: start}
            window = range(start, min(start + 10, len(texts)))
            for key, needles in required.items():
                if key == first_key:
                    continue
                # Les sous-chaînes sont essayées dans l'ordre de préférence :
                # "mensualite" avant "echeance", pour ne pas prendre la
                # colonne "No. d'échéance" pour celle des montants. Une même
                # cellule ne sert qu'une fois (un titre comme "Évolution
                # endettement et patrimoine" n'est pas un en-tête de tableau).
                used = set(colmap.values())
                for needle in needles:
                    match = next((c for c in window if c not in used and needle in texts[c]), None)
                    if match is not None:
                        colmap[key] = match
                        break
            if len(colmap) == len(required):
                found.append((r, colmap))
    return found


# ---------------------------------------------------------------------------
# Onglet "prêt" : tableau d'amortissement Date / Mensualité / Intérêts /
# Capital remboursé / Capital restant dû
# ---------------------------------------------------------------------------
def parse_loan_sheet(rows: list[list], name: str, today: date) -> dict | None:
    headers = _find_header_rows(rows, {
        "date": ("date",),
        "payment": ("mensualite", "echeance"),
        "interest": ("interet",),
        "remaining": ("restant",),
    })
    if not headers:
        return None
    header_row, cols = headers[0]
    capital_col = None
    for c, value in enumerate(rows[header_row]):
        if "rembours" in norm(value) and "capital" in norm(value):
            capital_col = c
            break

    schedule: list[dict] = []
    misses = 0
    for r in range(header_row + 1, len(rows)):
        d = _as_date(_cell(rows, r, cols["date"]))
        payment = _as_number(_cell(rows, r, cols["payment"]))
        interest = _as_number(_cell(rows, r, cols["interest"]))
        remaining = _as_number(_cell(rows, r, cols["remaining"]))
        if d is None or payment is None or remaining is None:
            misses += 1
            if misses >= 3 and schedule:
                break
            continue
        misses = 0
        capital = _as_number(_cell(rows, r, capital_col)) if capital_col is not None else None
        schedule.append({
            "date": d,
            "payment": payment,
            "interest": interest or 0.0,
            "capital": capital if capital is not None else payment - (interest or 0.0),
            "remaining": remaining,
        })
    if len(schedule) < 2:
        return None

    current_idx = None
    for i, row in enumerate(schedule):
        if row["date"] <= today:
            current_idx = i
        else:
            break
    last = schedule[-1]
    first = schedule[0]
    original_principal = first["remaining"] + first["capital"]
    remaining_today = schedule[current_idx]["remaining"] if current_idx is not None else original_principal
    finished = remaining_today <= 1 or last["date"] < date(today.year, today.month, 1)

    # L'app modélise un prêt à mensualité constante. Un vrai échéancier peut
    # commencer par un différé (intérêts seuls pendant des travaux) ou
    # contenir une pause de remboursement : on "recale" alors le prêt au
    # début de la dernière période à mensualité constante (celle en cours
    # aujourd'hui), pour que le capital restant dû calculé par l'app tombe
    # juste. Le capital emprunté d'origine est renvoyé à part, pour info.
    # Périodes à mensualité constante (à 0,5 % près), dans l'ordre.
    segments: list[tuple[int, int]] = []
    seg_start = 0
    for i in range(1, len(schedule) + 1):
        if i == len(schedule) or abs(schedule[i]["payment"] - schedule[seg_start]["payment"]) > 0.005 * max(
            schedule[seg_start]["payment"], 0.01
        ):
            segments.append((seg_start, i - 1))
            seg_start = i
    paying = [seg for seg in segments if schedule[seg[0]]["payment"] > 0]
    if not paying:
        return None
    if current_idx is not None and not finished:
        # Prêt en cours : la période qui contient aujourd'hui (ou, pendant
        # une pause sans mensualité, la dernière période payée avant).
        anchor_idx = next(
            (start for start, end in reversed(paying) if start <= current_idx), paying[0][0]
        )
    else:
        # Prêt pas encore commencé ou déjà terminé : la plus longue période
        # régulière (une dernière échéance d'ajustement ou un différé du
        # début ne doivent pas servir de référence).
        anchor_idx = max(paying, key=lambda seg: (seg[1] - seg[0], -seg[0]))[0]
    anchor = schedule[anchor_idx]
    principal = anchor["remaining"] + anchor["capital"]
    if principal <= 0 or anchor["payment"] <= 0:
        return None
    rate = anchor["interest"] / principal * 1200

    return {
        "property_name": name.strip()[:80],
        "principal_amount": round(principal, 2),
        "annual_rate_pct": round(rate, 3),
        "monthly_payment": round(anchor["payment"], 2),
        "start_date": anchor["date"].isoformat(),
        "end_date": last["date"].isoformat(),
        "installments": len(schedule),
        "remaining_today": round(max(remaining_today, 0.0), 2),
        "finished": finished,
        "original_principal": round(original_principal, 2),
        "original_start_date": first["date"].isoformat(),
        "rebased": anchor_idx > 0,
    }


# ---------------------------------------------------------------------------
# Onglet "endettement / patrimoine" :
#   - un tableau par bien (Mensualité, Capital restant dû, Loyer...) ;
#   - un historique daté Endettement / Patrimoine.
# Un tableau dont la colonne de libellés contient des mois ou des années
# (cumul mois par mois de tous les biens) est volontairement ignoré : il se
# recalcule à partir des prêts.
# ---------------------------------------------------------------------------
def _parse_property_table(rows: list[list]) -> list[dict]:
    properties: list[dict] = []
    for header_row, cols in _find_header_rows(rows, {
        "payment": ("mensualite",),
        "remaining": ("restant",),
        "rent": ("loyer",),
    }):
        label_col = cols["payment"] - 1
        if label_col < 0:
            continue
        block: list[dict] = []
        empties = 0
        is_timeline = False
        for r in range(header_row + 1, len(rows)):
            label = _cell(rows, r, label_col)
            if label is None or (isinstance(label, str) and not label.strip()):
                empties += 1
                if empties >= 3:
                    break
                continue
            empties = 0
            if _is_month_or_year_label(label):
                is_timeline = True
                break
            if not isinstance(label, str):
                break
            if norm(label).startswith("total"):
                break
            payment = _as_number(_cell(rows, r, cols["payment"]))
            remaining = _as_number(_cell(rows, r, cols["remaining"]))
            rent = _as_number(_cell(rows, r, cols["rent"]))
            block.append({
                "name": label.strip()[:80],
                "monthly_payment": round(payment, 2) if payment is not None else None,
                "remaining": round(remaining, 2) if remaining is not None else None,
                "monthly_rent": round(rent, 2) if rent is not None else None,
                # Ligne présente mais sans mensualité : prêt soldé ou bien vendu.
                "loan_running": payment is not None and payment > 0,
            })
        if block and not is_timeline:
            properties.extend(block)
    return properties


def _parse_net_worth_history(rows: list[list]) -> list[dict]:
    snapshots: list[dict] = []
    seen_dates: set[str] = set()
    for header_row, cols in _find_header_rows(rows, {
        "debts": ("endettement",),
        "assets": ("patrimoine",),
    }):
        # Colonne des dates : la plus proche à gauche de "Endettement" qui
        # contient des dates juste en dessous de l'en-tête.
        date_col = None
        for c in range(cols["debts"] - 1, -1, -1):
            if any(_as_date(_cell(rows, r, c)) for r in range(header_row + 1, min(header_row + 6, len(rows)))):
                date_col = c
                break
        if date_col is None:
            continue
        misses = 0
        for r in range(header_row + 1, len(rows)):
            d = _as_date(_cell(rows, r, date_col))
            debts = _as_number(_cell(rows, r, cols["debts"]))
            assets = _as_number(_cell(rows, r, cols["assets"]))
            if d is None or debts is None or assets is None:
                misses += 1
                if misses >= 5:
                    break
                continue
            misses = 0
            if debts < 0 or assets < 0 or d.isoformat() in seen_dates:
                continue
            seen_dates.add(d.isoformat())
            snapshots.append({"snapshot_date": d.isoformat(), "debts": round(debts, 2), "assets": round(assets, 2)})
    snapshots.sort(key=lambda s: s["snapshot_date"])
    return snapshots


def parse_net_worth_sheet(rows: list[list]) -> dict | None:
    properties = _parse_property_table(rows)
    snapshots = _parse_net_worth_history(rows)
    if not properties and not snapshots:
        return None
    return {"properties": properties, "snapshots": snapshots}


# ---------------------------------------------------------------------------
# Classeur complet
# ---------------------------------------------------------------------------
def analyze_workbook(raw_bytes: bytes, today: date) -> list[dict]:
    """Une entrée par onglet : son type probable, ce qui a pu en être lu,
    et s'il doit être coché par défaut."""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(raw_bytes), data_only=True, read_only=True)
    try:
        sheets: list[dict] = []
        for ws in wb.worksheets[:MAX_SHEETS]:
            rows = _grid(ws)
            filled = sum(1 for row in rows for v in row if v not in (None, ""))
            info = {
                "name": ws.title,
                "hidden": ws.sheet_state != "visible",
                "filled_cells": filled,
                "month": sheet_month(ws.title),
                "loan": None,
                "net_worth": None,
                "kind": "ignore",
                "selected": False,
                "note": None,
            }
            if filled == 0:
                info["note"] = "Onglet vide"
                sheets.append(info)
                continue
            loan = parse_loan_sheet(rows, ws.title, today)
            if loan:
                info.update(kind="loan", loan=loan, selected=True)
                sheets.append(info)
                continue
            net_worth = parse_net_worth_sheet(rows)
            if net_worth:
                info.update(kind="net_worth", net_worth=net_worth, selected=True)
                sheets.append(info)
                continue
            if info["month"]:
                info.update(kind="monthly", selected=True)
            sheets.append(info)
    finally:
        wb.close()

    # Rapprochement prêt <-> ligne du tableau par bien : loyer mensuel, et
    # prêt signalé comme soldé/bien vendu quand sa ligne n'a plus de
    # mensualité (bien vendu, prêt remboursé par anticipation...).
    by_key: dict[str, dict] = {}
    for sheet in sheets:
        if sheet["net_worth"]:
            for prop in sheet["net_worth"]["properties"]:
                by_key.setdefault(property_key(prop["name"]), prop)
    for sheet in sheets:
        loan = sheet["loan"]
        if not loan:
            continue
        prop = by_key.get(property_key(loan["property_name"]))
        loan["monthly_rent"] = prop["monthly_rent"] if prop else None
        if prop and not prop["loan_running"]:
            loan["finished"] = True
            sheet["note"] = "Plus de mensualité dans le tableau d'endettement : prêt soldé ou bien vendu"
        if loan["finished"]:
            sheet["selected"] = False
            if not sheet["note"]:
                sheet["note"] = "Prêt terminé"
    return sheets
